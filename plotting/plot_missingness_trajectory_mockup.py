#!/usr/bin/env python3
"""One-dataset IEMOCAP accuracy/compute trajectory mock-up from reported values.

Use a downloaded XLSX snapshot as --input. Each path joins one method's
reported means at 0, 10, 20, 30, and 40 percent modality missingness.
No missing values are interpolated. Error bars are omitted for this layout
prototype; parsed dispersion values remain available in the JSON sidecar.
Each run updates the same PNG, SVG, PDF, and data files using the default stem.
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.colors as mcolors
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
from matplotlib.lines import Line2D
from matplotlib.ticker import FixedLocator, FuncFormatter, NullLocator


SOURCE_URL = (
    "https://docs.google.com/spreadsheets/d/"
    "1ZsKgZztYMPOj76LL7eYWT07JJ4k0cWdaC1zmOlIy3q0/edit"
)
SHEET = "Main_ResultsFull"
METHOD_ROWS = {
    "MBT": "Vanilla MBT (R)",
    "MAESTRO": "MAESTRO (R)",
    "ShaSpec": "ShaSpec (R)",
    "DecALign": "DecALign (R)",
    "MultiModN": "MultiModN (R)",
    "DyMo": "DyMo (E)",
    "AdaMML": "AdaMML (E)",
    "DyMM": "DyMM (E)",
    "Greedy Submodular": "Greedy Submodular (own backbone+selection)",
    "JAFA": "JAFA (own backbone+selection, r_cost=0.05)",
    "MMEE": "MMEE",
    "SeMARC w/o RL": "Ours w/o RL",
    "SeMARC": "Ours",
}
MISSINGNESS = [0, 10, 20, 30, 40]
OURS_COLOR = "#963E92"
ABLATION_COLOR = "#9C8ABB"
BASELINE_COLOR = "#6D777E"
DISPLAY_LABELS = {
    "SeMARC w/o RL": "SeMA",
    "Greedy Submodular": "GMS",
}


def parse_f1(value):
    parts = str(value).split("|")
    return float(parts[0].strip()), float(parts[1].strip()) if len(parts) > 1 else None


def read_iemocap(path: Path):
    book = openpyxl.load_workbook(path, data_only=True)
    sheet = book[SHEET]
    start = next(
        c for c in range(1, sheet.max_column + 1)
        if sheet.cell(1, c).value == "IEMOCAP"
    )
    rows = {
        str(sheet.cell(r, 1).value).strip(): r
        for r in range(1, sheet.max_row + 1)
    }
    data = {}
    for name, source_name in METHOD_ROWS.items():
        row = rows[source_name]
        values = {"f1": [], "dispersion": [], "gflops": [], "mu": [], "source_cells": []}
        for i, rate in enumerate(MISSINGNESS):
            col = start + 3 * i
            observed_rate = float(sheet.cell(2, col).value) * 100
            if abs(observed_rate - rate) > 1e-8:
                raise ValueError(f"Missingness header mismatch at column {col}")
            if [sheet.cell(3, col + j).value for j in range(3)] != ["F1", "MU", "GFLOP"]:
                raise ValueError(f"Metric header mismatch at column {col}")
            f1, dispersion = parse_f1(sheet.cell(row, col).value)
            compute = float(sheet.cell(row, col + 2).value)
            if not 0 <= f1 <= 1 or not np.isfinite(compute) or compute <= 0:
                raise ValueError(f"Invalid measurement for {name} at {rate}%")
            values["f1"].append(f1)
            values["dispersion"].append(dispersion)
            values["gflops"].append(compute)
            values["mu"].append(float(sheet.cell(row, col + 1).value))
            values["source_cells"].append(
                f"{sheet.cell(row, col).coordinate}:{sheet.cell(row, col + 2).coordinate}"
            )
        data[name] = values
    book.close()
    return data


def shaded(color, fraction):
    return tuple((1 - fraction) + fraction * np.array(mcolors.to_rgb(color)))


def render(data, output_dir: Path, stem: str):
    plt.rcParams.update({
        "font.family": "DejaVu Sans",
        "font.size": 18.5,
        "axes.labelsize": 24,
        "xtick.labelsize": 18,
        "ytick.labelsize": 18,
        "axes.linewidth": 1.0,
        "axes.edgecolor": "black",
        "pdf.fonttype": 42,
        "ps.fonttype": 42,
        "svg.fonttype": "none",
        "savefig.facecolor": "white",
    })
    fig, ax = plt.subplots(figsize=(11.2, 7.0))
    fig.subplots_adjust(left=.115, right=.975, top=.86, bottom=.14)
    ax.set_xscale("log")
    ax.set_xlim(.079, 17)
    ax.set_ylim(.325, .726)
    ax.xaxis.set_major_locator(FixedLocator([.1, .2, .5, 1, 2, 5, 10]))
    ax.xaxis.set_major_formatter(FuncFormatter(lambda value, position: f"{value:g}"))
    ax.xaxis.set_minor_locator(NullLocator())
    ax.yaxis.set_major_locator(FixedLocator(np.arange(.35, .71, .05)))
    ax.grid(color="#E5E7E9", linewidth=.65)
    ax.set_axisbelow(True)
    ax.set_xlabel("GFLOPs per inference (log scale)  ·  lower is better", labelpad=7)
    ax.set_ylabel("Macro-F1  ·  higher is better", labelpad=7)
    for spine in ax.spines.values():
        spine.set_visible(True)
        spine.set_color("black")
        spine.set_linewidth(1.0)
    ax.tick_params(length=3, color="black", pad=4)

    for name, values in data.items():
        ours = name == "SeMARC"
        ablation = name == "SeMARC w/o RL"
        color = OURS_COLOR if ours else ABLATION_COLOR if ablation else BASELINE_COLOR
        linewidth = 2.8 if ours else 1.9 if ablation else 1.05
        zorder = 8 if ours else 5 if ablation else 3
        ax.plot(
            values["gflops"], values["f1"], color=color,
            linewidth=linewidth, linestyle="--" if ablation else "-",
            alpha=1 if ours or ablation else .72, zorder=zorder,
        )
        for i, (compute, f1) in enumerate(zip(values["gflops"], values["f1"])):
            ax.plot(
                compute, f1, linestyle="none", marker="o",
                markersize=7.4 if ours else 5.8 if ablation else 4.8,
                markerfacecolor=shaded(color, i / 4),
                markeredgecolor=color, markeredgewidth=1.1 if ours else .8,
                zorder=zorder + 1,
            )

    # Each label points to the named method, at its 0% or 40% observation.
    labels = {
        "MBT": (0, (.50, .651)),
        "MAESTRO": (4, (.275, .359)),
        "ShaSpec": (0, (1.04, .638)),
        "DecALign": (0, (8.0, .506)),
        "MultiModN": (0, (.73, .607)),
        "DyMo": (4, (.56, .463)),
        "AdaMML": (4, (.251, .480)),
        "DyMM": (0, (2.04, .625)),
        "Greedy Submodular": (4, (.091, .408)),
        "JAFA": (4, (.087, .465)),
        "MMEE": (4, (.51, .518)),
        "SeMARC w/o RL": (0, (.51, .699)),
        "SeMARC": (0, (.175, .711)),
    }
    for name, (index, destination) in labels.items():
        values = data[name]
        ours = name == "SeMARC"
        ablation = name == "SeMARC w/o RL"
        color = OURS_COLOR if ours else ABLATION_COLOR if ablation else "#454D53"
        ax.annotate(
            DISPLAY_LABELS.get(name, name),
            xy=(values["gflops"][index], values["f1"][index]),
            xytext=destination, textcoords="data", fontsize=18 if not ours else 20,
            fontweight="bold" if ours else "normal", color=color,
            ha="left", va="center",
            arrowprops={"arrowstyle": "-", "color": color, "lw": .65, "shrinkB": 5},
            zorder=20,
        )

    ours = data["SeMARC"]
    rate_offsets = [(-12, 5), (-12, 5), (-12, 9), (-12, -4), (10, -5)]
    for i, rate in enumerate(MISSINGNESS):
        ax.annotate(
            f"{rate}%", xy=(ours["gflops"][i], ours["f1"][i]),
            xytext=rate_offsets[i], textcoords="offset points",
            fontsize=17, ha="left" if i == 4 else "right",
            va="center", color=OURS_COLOR, zorder=20,
        )
    ax.annotate(
        "", xy=(ours["gflops"][-1], ours["f1"][-1]),
        xytext=(ours["gflops"][-2], ours["f1"][-2]),
        arrowprops={"arrowstyle": "-|>", "color": OURS_COLOR, "lw": 2.3,
                    "shrinkA": 6, "shrinkB": 6, "mutation_scale": 13}, zorder=10,
    )

    # One shared fill scale encodes missingness for every method.
    handles = [Line2D([], [], marker="o", linestyle="none", markersize=6,
                      markeredgecolor=BASELINE_COLOR,
                      markerfacecolor=shaded(BASELINE_COLOR, i / 4), label=f"{rate}%")
               for i, rate in enumerate(MISSINGNESS)]
    # A text-only first entry keeps the legend heading on the same row.
    handles.insert(0, Line2D([], [], linestyle="none", label="Modality missingness:"))
    legend = ax.legend(
        handles=handles, ncol=len(handles), frameon=True,
        fancybox=False, framealpha=1.0, facecolor="white", edgecolor="#B9B9B9",
        loc="lower center", bbox_to_anchor=(.5, 1.013), fontsize=17.5,
        borderaxespad=0, borderpad=.25,
        columnspacing=.85, handletextpad=.45, handlelength=1,
    )
    legend.get_texts()[0].set_fontsize(18)
    legend.get_frame().set_linewidth(.6)
    output_dir.mkdir(parents=True, exist_ok=True)
    output_paths = []
    for extension in ("png", "svg", "pdf"):
        path = output_dir / f"{stem}.{extension}"
        fig.savefig(path, dpi=400, bbox_inches="tight", pad_inches=.05)
        output_paths.append(path)
    plt.close(fig)
    return output_paths


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output-dir", type=Path, default=Path(__file__).parent / "output")
    parser.add_argument("--stem", default="iemocap_missingness_trajectory_publication")
    args = parser.parse_args()
    data = read_iemocap(args.input)
    paths = render(data, args.output_dir, args.stem)
    sidecar = args.output_dir / f"{args.stem}_data.json"
    sidecar.write_text(json.dumps({
        "source_url": SOURCE_URL,
        "source_workbook": str(args.input.resolve()),
        "sheet": SHEET,
        "dataset": "IEMOCAP",
        "missingness_percent": MISSINGNESS,
        "methods": data,
    }, indent=2) + "\n", encoding="utf-8")
    for path in [*paths, sidecar]:
        print(path.resolve())


if __name__ == "__main__":
    main()
