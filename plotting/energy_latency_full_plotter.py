#!/usr/bin/env python3
"""Plot energy/latency measurements for SeMARC and four comparison models.

By default, one selected dataset is shown in separate GPU, CPU, and Android
panels with linear axes. Passing ``--dataset all`` produces the complementary
dataset grid, where platforms share each dataset's log-log panel. Model identity
is encoded consistently by color and marker shape. The selected-dataset view
labels the full-availability F1 means from ``Main_ResultsFull``; these are shared
reference scores, not separate hardware-specific accuracy measurements. It omits
the SeMARC-without-RL ablation and enlarges baseline clusters in inset axes.
Empty spreadsheet blocks are reported without filling missing measurements.

Raw latency is recorded in milliseconds. GPU and Android energy are recorded in
millijoules, while CPU energy is recorded in joules; all plotted values are
converted to seconds and joules.

Dependencies: matplotlib, openpyxl
"""

from __future__ import annotations

import argparse
import math
import re
import urllib.request
from dataclasses import dataclass
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import openpyxl
from matplotlib.lines import Line2D
from matplotlib.patches import ConnectionPatch, Rectangle
from matplotlib.ticker import FuncFormatter, LogLocator, MaxNLocator, NullFormatter


SHEET_URL = (
    "https://docs.google.com/spreadsheets/d/"
    "1ZsKgZztYMPOj76LL7eYWT07JJ4k0cWdaC1zmOlIy3q0/"
    "export?format=xlsx"
)

METHOD_ORDER = [
    "SeMARC",
    "MBT",
    "AdaMML",
    "DyMM",
    "DyMo",
]
METHOD_LABELS = {
    "SeMARC": "SeMARC (Ours)",
    "SeMARC w/o RL": "SeMA",
    "MBT": "MBT (Monolithic Fusion)",
    "AdaMML": "AdaMML (Selective Fusion)",
    "DyMM": "DyMM (Selective Fusion)",
    "DyMo": "DyMo (Selective Fusion)",
}
METHOD_COLORS = {
    "SeMARC": "#AA4499",  # Paul Tol purple
    "SeMARC w/o RL": "#CC6677",  # Paul Tol rose
    "MBT": "#EE7733",  # Paul Tol orange
    "AdaMML": "#4477AA",  # Paul Tol blue
    "DyMM": "#228833",  # Paul Tol green
    "DyMo": "#882255",  # Paul Tol wine
}
RESULT_ROW_LABELS = {
    "SeMARC": "Ours",
    "MBT": "Vanilla MBT (R)",
    "AdaMML": "AdaMML (E)",
    "DyMM": "DyMM (E)",
    "DyMo": "DyMo (E)",
}

DATASET_ORDER = [
    "IEMOCAP",
    "MM-Fi",
    "CMI",
    "CZU-MHAD",
    "DSADS",
    "EAV",
    "UTD-MHAD",
]
METHOD_MARKERS = {
    "SeMARC": "*",
    "SeMARC w/o RL": "P",
    "MBT": "o",
    "AdaMML": "s",
    "DyMM": "D",
    "DyMo": "p",
}

PLATFORM_ANCHORS = {"GPU", "CPU", "Onnx_Android"}
PLATFORM_ORDER = [
    "GPU",
    "CPU",
    "Android (Full model)",
    "Android (FP16)",
    "Android (INT8)",
]
PLATFORM_FILLSTYLES = {
    "GPU": "full",
    "CPU": "none",
    "Android (Full model)": "left",
    "Android (FP16)": "right",
    "Android (INT8)": "bottom",
}
PLATFORM_LABELS = {
    "GPU": "GPU (filled)",
    "CPU": "CPU (hollow)",
    "Android (Full model)": "Android full model",
    "Android (FP16)": "Android FP16",
    "Android (INT8)": "Android INT8 (half-filled)",
}


@dataclass(frozen=True)
class Measurement:
    platform: str
    dataset: str
    method: str
    latency_s: float
    energy_j: float
    source_row: int


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Plot all populated Energy_Study latency/energy measurements."
    )
    parser.add_argument(
        "--input",
        default=SHEET_URL,
        help="Local XLSX path or Google Sheets XLSX export URL.",
    )
    parser.add_argument("--sheet", default="Energy_Study")
    parser.add_argument("--results-sheet", default="Main_ResultsFull")
    parser.add_argument(
        "--dataset",
        default="IEMOCAP",
        help="Dataset for platform panels; use 'all' for the seven-dataset grid.",
    )
    parser.add_argument(
        "--scale",
        choices=("linear", "log"),
        default="linear",
        help="Axis scale for a selected dataset (default: linear).",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=Path(__file__).resolve().parent / "output",
    )
    parser.add_argument("--stem", default=None)
    parser.add_argument("--dpi", type=int, default=400)
    return parser.parse_args()


def load_workbook(source: str):
    if source.startswith(("http://", "https://")):
        request = urllib.request.Request(
            source, headers={"User-Agent": "Mozilla/5.0"}
        )
        with urllib.request.urlopen(request, timeout=60) as response:
            return openpyxl.load_workbook(
                BytesIO(response.read()), data_only=True
            )
    return openpyxl.load_workbook(
        Path(source).expanduser(), data_only=True
    )


def first_number(value) -> float:
    """Return the first numeric value from a number or a pipe-delimited cell."""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    text = str(value).split("|")[0].strip().replace(",", "")
    match = re.search(r"[-+]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][-+]?\d+)?", text)
    if match is None:
        raise ValueError(f"Expected a numeric measurement, got {value!r}")
    return float(match.group(0))


def canonical_method(raw: object) -> str:
    name = " ".join(str(raw or "").split())
    if name == "Ours":
        return "SeMARC"
    if name.startswith("Ours w/o RL"):
        return "SeMARC w/o RL"
    if name.startswith("Vanilla MBT"):
        return "MBT"
    return name


def platform_name(anchor: str, variant: object) -> str:
    if anchor != "Onnx_Android":
        return anchor
    variant_name = " ".join(str(variant or "").split())
    return {
        "Full Model": "Android (Full model)",
        "Quantized (FP16)": "Android (FP16)",
        "Quantized (INT8)": "Android (INT8)",
    }.get(variant_name, f"Android ({variant_name or 'unspecified'})")


def energy_scale(anchor: str, metric_headers: list[str]) -> float:
    explicit = " ".join(metric_headers).lower()
    if "energy(mj)" in explicit or "energy (mj)" in explicit:
        return 1e-3
    if "energy(j)" in explicit or "energy (j)" in explicit:
        return 1.0
    # The sheet's unlabeled energy columns inherit the unit from their block.
    return 1.0 if anchor == "CPU" else 1e-3


def discover_blocks(ws) -> list[tuple[int, int]]:
    anchors = [
        row
        for row in range(1, ws.max_row + 1)
        if str(ws.cell(row, 1).value or "").strip() in PLATFORM_ANCHORS
    ]
    return [
        (row, anchors[index + 1] - 1 if index + 1 < len(anchors) else ws.max_row)
        for index, row in enumerate(anchors)
    ]


def discover_method_columns(ws, method_row: int, metric_row: int):
    """Return (method, latency_column, energy_column) entries for one block."""
    latency_columns = [
        column
        for column in range(1, ws.max_column + 1)
        if str(ws.cell(metric_row, column).value or "")
        .strip()
        .lower()
        .startswith("latency")
    ]
    used_latency_columns: set[int] = set()
    entries = []

    for energy_column in range(1, ws.max_column + 1):
        metric = str(ws.cell(metric_row, energy_column).value or "").strip()
        raw_method = ws.cell(method_row, energy_column).value
        if not metric.lower().startswith("energy") or raw_method is None:
            continue

        candidates = [
            column
            for column in latency_columns
            if column not in used_latency_columns and abs(column - energy_column) <= 3
        ]
        if not candidates:
            raise ValueError(
                f"No latency column found near {ws.cell(metric_row, energy_column).coordinate}"
            )
        latency_column = min(candidates, key=lambda column: abs(column - energy_column))
        used_latency_columns.add(latency_column)
        entries.append(
            (canonical_method(raw_method), latency_column, energy_column)
        )

    return entries


def read_measurements(book, sheet_name: str):
    if sheet_name not in book.sheetnames:
        raise KeyError(
            f"Worksheet {sheet_name!r} not found; available: {book.sheetnames}"
        )
    ws = book[sheet_name]
    measurements: list[Measurement] = []
    empty_blocks: list[str] = []
    incomplete_pairs: list[str] = []

    for anchor_row, block_end in discover_blocks(ws):
        anchor = str(ws.cell(anchor_row, 1).value).strip()
        platform = platform_name(anchor, ws.cell(anchor_row, 2).value)
        method_row = anchor_row + 1
        metric_row = anchor_row + 2
        method_columns = discover_method_columns(ws, method_row, metric_row)
        headers = [
            str(ws.cell(metric_row, energy_column).value or "")
            for _, _, energy_column in method_columns
        ]
        e_scale = energy_scale(anchor, headers)
        block_count = 0

        for row in range(metric_row + 1, block_end + 1):
            dataset = " ".join(str(ws.cell(row, 1).value or "").split())
            if not dataset:
                continue
            for method, latency_column, energy_column in method_columns:
                raw_latency = ws.cell(row, latency_column).value
                raw_energy = ws.cell(row, energy_column).value
                if raw_latency in (None, "") and raw_energy in (None, ""):
                    continue
                if raw_latency in (None, "") or raw_energy in (None, ""):
                    incomplete_pairs.append(
                        f"Incomplete latency/energy pair for {platform}, {dataset}, "
                        f"{method} at row {row}"
                    )
                    continue
                latency_s = first_number(raw_latency) * 1e-3
                energy_j = first_number(raw_energy) * e_scale
                if latency_s <= 0 or energy_j <= 0:
                    raise ValueError(
                        f"Measurements must be positive; got {latency_s=}, "
                        f"{energy_j=} at row {row}"
                    )
                measurements.append(
                    Measurement(
                        platform=platform,
                        dataset=dataset,
                        method=method,
                        latency_s=latency_s,
                        energy_j=energy_j,
                        source_row=row,
                    )
                )
                block_count += 1

        if block_count == 0:
            empty_blocks.append(platform)

    if not measurements:
        raise ValueError(f"No populated latency/energy pairs found in {sheet_name}")

    keys = [(m.platform, m.dataset, m.method) for m in measurements]
    if len(keys) != len(set(keys)):
        raise ValueError("Duplicate platform/dataset/method measurements found")
    return measurements, empty_blocks, incomplete_pairs


def read_f1_scores(book, sheet_name: str, dataset: str) -> dict[str, float]:
    """Read full-availability F1 means for the displayed methods."""
    if sheet_name not in book.sheetnames:
        raise KeyError(
            f"Worksheet {sheet_name!r} not found; available: {book.sheetnames}"
        )
    ws = book[sheet_name]
    dataset_column = next(
        (
            column
            for column in range(1, ws.max_column + 1)
            if str(ws.cell(1, column).value or "").strip().casefold()
            == dataset.casefold()
        ),
        None,
    )
    if dataset_column is None:
        raise KeyError(f"No {dataset!r} block found in {sheet_name}")
    if first_number(ws.cell(2, dataset_column).value) != 0.0:
        raise ValueError(
            f"Expected zero missingness at {ws.cell(2, dataset_column).coordinate}"
        )
    if str(ws.cell(3, dataset_column).value or "").strip() != "F1":
        raise ValueError(
            f"Expected F1 header at {ws.cell(3, dataset_column).coordinate}"
        )

    row_by_label = {
        str(ws.cell(row, 1).value or "").strip(): row
        for row in range(1, ws.max_row + 1)
    }
    scores = {}
    for method in METHOD_ORDER:
        source_label = RESULT_ROW_LABELS[method]
        if source_label not in row_by_label:
            raise KeyError(f"Missing result row {source_label!r} in {sheet_name}")
        scores[method] = first_number(
            ws.cell(row_by_label[source_label], dataset_column).value
        )
    return scores


def ordered_present(values: set[str], preferred: list[str]) -> list[str]:
    return [value for value in preferred if value in values] + sorted(
        values - set(preferred)
    )


FONT_SIZE_INCREMENT = 4.0
MAIN_MARKER_AREA_SCALE = 8.0 / 5.0


def publication_font(previous_size: float) -> float:
    """Apply the requested four-point increase consistently to every label."""
    return previous_size + FONT_SIZE_INCREMENT


def configure_style() -> None:
    plt.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "font.size": publication_font(11.0),
            "axes.labelsize": publication_font(13.0),
            "xtick.labelsize": publication_font(10.5),
            "ytick.labelsize": publication_font(10.5),
            "legend.fontsize": publication_font(9.5),
            "figure.labelsize": publication_font(13.2),
            "axes.edgecolor": "black",
            "axes.linewidth": 1.0,
            "pdf.fonttype": 42,
            "ps.fonttype": 42,
            "svg.fonttype": "none",
        }
    )


def save_figure(fig, output_dir: Path, stem: str, dpi: int) -> list[Path]:
    """Export tightly cropped vector and raster versions without an outer frame."""
    fig.canvas.draw()
    content_bounds = fig.get_tightbbox(fig.canvas.get_renderer())
    output_dir.mkdir(parents=True, exist_ok=True)
    paths = []
    for suffix, options in (
        ("pdf", {}),
        ("svg", {}),
        ("png", {"dpi": dpi}),
    ):
        path = output_dir / f"{stem}.{suffix}"
        fig.savefig(
            path, bbox_inches=content_bounds.padded(0.060),
            facecolor="white", **options,
        )
        paths.append(path)
    plt.close(fig)
    return paths


def compact_tick(value: float, _position: int) -> str:
    """Format narrow-range log ticks without scientific-notation clutter."""
    if value >= 1000:
        return f"{value / 1000:g}k"
    if value >= 1:
        return f"{value:g}"
    return f"{value:.3g}"


def f1_label(score: float) -> str:
    """Round the recorded decimal mean conventionally (0.595 becomes 0.60)."""
    rounded = Decimal(str(score)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    return f"F1={rounded:.2f}"


def render(
    measurements: list[Measurement],
    output_dir: Path,
    stem: str,
    dpi: int,
) -> list[Path]:
    configure_style()
    plot_measurements = [
        measurement
        for measurement in measurements
        if measurement.method in METHOD_ORDER
    ]
    methods = ordered_present(
        {measurement.method for measurement in plot_measurements}, METHOD_ORDER
    )
    datasets = ordered_present(
        {measurement.dataset for measurement in plot_measurements}, DATASET_ORDER
    )
    platforms = ordered_present(
        {measurement.platform for measurement in plot_measurements}, PLATFORM_ORDER
    )

    columns = min(4, len(datasets) + 1)
    rows = math.ceil((len(datasets) + 1) / columns)
    fig, axes = plt.subplots(
        rows,
        columns,
        figsize=(3.2 * columns, 3.2 * rows),
        sharex=False,
        sharey=False,
        squeeze=False,
    )
    flat_axes = list(axes.flat)

    for panel_index, (ax, dataset) in enumerate(zip(flat_axes, datasets)):
        panel_values = [
            measurement
            for measurement in plot_measurements
            if measurement.dataset == dataset
        ]
        for measurement in panel_values:
            ours = measurement.method == "SeMARC"
            color = METHOD_COLORS.get(measurement.method, "#666666")
            fillstyle = PLATFORM_FILLSTYLES.get(
                measurement.platform, "full"
            )
            marker_edge = (
                "white" if measurement.platform == "GPU" else color
            )
            ax.plot(
                measurement.latency_s,
                measurement.energy_j,
                linestyle="none",
                marker=METHOD_MARKERS.get(measurement.method, "o"),
                markersize=(
                    14.5
                    if ours
                    else 9.5
                    if measurement.method == "SeMARC w/o RL"
                    else 8.0
                ),
                markerfacecolor=color,
                markerfacecoloralt="white",
                markeredgecolor=marker_edge,
                markeredgewidth=1.5 if ours else 1.2,
                fillstyle=fillstyle,
                alpha=1.0 if ours else 0.94,
                zorder=7 if ours else 5,
            )

        ax.set_xscale("log")
        ax.set_yscale("log")
        ax.xaxis.set_major_locator(LogLocator(base=10, numticks=7))
        ax.yaxis.set_major_locator(LogLocator(base=10, numticks=7))
        ax.xaxis.set_minor_locator(
            LogLocator(base=10, subs=(2, 5), numticks=14)
        )
        ax.yaxis.set_minor_locator(
            LogLocator(base=10, subs=(2, 5), numticks=14)
        )
        ax.xaxis.set_minor_formatter(NullFormatter())
        ax.yaxis.set_minor_formatter(NullFormatter())
        x_values = [measurement.latency_s for measurement in panel_values]
        y_values = [measurement.energy_j for measurement in panel_values]
        x_log = [math.log10(value) for value in x_values]
        y_log = [math.log10(value) for value in y_values]
        x_pad = 0.16 * max(max(x_log) - min(x_log), 0.55)
        y_pad = 0.16 * max(max(y_log) - min(y_log), 0.55)
        ax.set_xlim(10 ** (min(x_log) - x_pad), 10 ** (max(x_log) + x_pad))
        ax.set_ylim(10 ** (min(y_log) - y_pad), 10 ** (max(y_log) + y_pad))
        if max(x_values) / min(x_values) < 10:
            ax.xaxis.set_major_locator(
                LogLocator(base=10, subs=(1, 2, 3, 5), numticks=12)
            )
            ax.xaxis.set_major_formatter(FuncFormatter(compact_tick))
        if max(y_values) / min(y_values) < 10:
            ax.yaxis.set_major_locator(
                LogLocator(base=10, subs=(1, 2, 3, 5), numticks=12)
            )
            ax.yaxis.set_major_formatter(FuncFormatter(compact_tick))
        ax.grid(
            True,
            which="major",
            color="#D8D8D8",
            linewidth=0.55,
            alpha=0.82,
        )
        ax.grid(
            True,
            which="minor",
            color="#ECECEC",
            linewidth=0.35,
            alpha=0.62,
        )
        ax.tick_params(
            which="major", direction="out", length=3.0, width=0.65, pad=2.0
        )
        ax.tick_params(which="minor", direction="out", length=1.8, width=0.45)
        ax.text(
            0.035,
            0.955,
            f"({chr(ord('a') + panel_index)}) {dataset}",
            transform=ax.transAxes,
            ha="left",
            va="top",
            fontsize=publication_font(11.2),
            color="#333333",
        )
        for spine in ax.spines.values():
            spine.set_color("black")
            spine.set_linewidth(1.0)

    method_handles = [
        Line2D(
            [0],
            [0],
            linestyle="none",
            marker=METHOD_MARKERS.get(method, "o"),
            markerfacecolor=METHOD_COLORS.get(method, "#666666"),
            markerfacecoloralt="white",
            markeredgecolor="white",
            markeredgewidth=0.8,
            markersize=11.5 if method == "SeMARC" else 7.5,
            fillstyle="full",
            label=METHOD_LABELS.get(method, method),
        )
        for method in methods
    ]
    platform_handles = [
        Line2D(
            [0],
            [0],
            linestyle="none",
            marker="o",
            markerfacecolor="#666666",
            markerfacecoloralt="white",
            markeredgecolor=(
                "white" if platform == "GPU" else "#666666"
            ),
            markeredgewidth=1.1,
            markersize=7.5,
            fillstyle=PLATFORM_FILLSTYLES.get(platform, "full"),
            label=PLATFORM_LABELS.get(platform, platform),
        )
        for platform in platforms
    ]

    legend_ax = flat_axes[len(datasets)]
    legend_ax.set_axis_off()
    method_legend = legend_ax.legend(
        handles=method_handles,
        loc="upper left",
        bbox_to_anchor=(0.0, 1.0),
        ncol=1,
        title="Model",
        frameon=True,
        fancybox=False,
        framealpha=1.0,
        facecolor="#FAFAFA",
        edgecolor="black",
        columnspacing=0.6,
        handlelength=0.9,
        handletextpad=0.35,
        borderpad=0.3,
        labelspacing=0.28,
    )
    method_legend.get_frame().set_linewidth(1.0)
    method_legend.get_title().set_fontsize(publication_font(9.5))

    platform_legend = legend_ax.legend(
        handles=platform_handles,
        loc="lower left",
        bbox_to_anchor=(0.0, 0.0),
        ncol=1,
        title="Compute platform",
        frameon=True,
        fancybox=False,
        framealpha=1.0,
        facecolor="#FAFAFA",
        edgecolor="black",
        columnspacing=0.6,
        handlelength=0.8,
        handletextpad=0.35,
        borderpad=0.3,
        labelspacing=0.28,
    )
    platform_legend.get_frame().set_linewidth(1.0)
    platform_legend.get_title().set_fontsize(publication_font(9.5))
    legend_ax.add_artist(method_legend)
    legend_ax.text(
        0.02,
        0.39,
        "CPU values for MBT\nare estimated.",
        transform=legend_ax.transAxes,
        ha="left",
        va="center",
        fontsize=publication_font(8.2),
        color="#555555",
    )

    for ax in flat_axes[len(datasets) + 1 :]:
        ax.set_visible(False)

    fig.supxlabel("Latency per inference (s, log scale)", y=0.018)
    fig.supylabel("Energy per inference (J, log scale)", x=0.012)
    fig.subplots_adjust(
        left=0.075,
        right=0.995,
        bottom=0.09,
        top=0.99,
        wspace=0.28,
        hspace=0.25,
    )

    return save_figure(fig, output_dir, stem, dpi)


def render_dataset_platforms(
    measurements: list[Measurement],
    dataset: str,
    scale: str,
    f1_scores: dict[str, float],
    output_dir: Path,
    stem: str,
    dpi: int,
) -> list[Path]:
    """Render one platform panel per dataset, with baseline zoom insets."""
    configure_style()
    panel_measurements = [
        measurement
        for measurement in measurements
        if measurement.dataset.casefold() == dataset.casefold()
        and measurement.method in METHOD_ORDER
    ]
    if not panel_measurements:
        available = ordered_present(
            {measurement.dataset for measurement in measurements}, DATASET_ORDER
        )
        raise ValueError(
            f"Dataset {dataset!r} has no measurements; available: {available}"
        )

    dataset_name = panel_measurements[0].dataset
    platforms = ordered_present(
        {measurement.platform for measurement in panel_measurements},
        PLATFORM_ORDER,
    )
    methods = ordered_present(
        {measurement.method for measurement in panel_measurements}, METHOD_ORDER
    )
    fig, axes = plt.subplots(
        1,
        len(platforms),
        figsize=(4.10 * len(platforms), 3.90),
        squeeze=False,
    )

    # Fixed inset-relative label anchors keep the enlarged text inside the box.
    annotation_positions = {
        "GPU": {
            "MBT": (0.26, 0.90),
            "AdaMML": (0.25, 0.10),
            "DyMM": (0.76, 0.12),
            "DyMo": (0.70, 0.50),
        },
        "CPU": {
            "DyMM": (0.26, 0.90),
            "DyMo": (0.76, 0.90),
            "MBT": (0.60, 0.55),
            "AdaMML": (0.62, 0.10),
        },
        "Android (INT8)": {
            "DyMo": (0.53, 0.90),
            "MBT": (0.24, 0.43),
            "DyMM": (0.75, 0.73),
            "AdaMML": (0.65, 0.10),
        },
    }

    def draw_point(ax, measurement: Measurement, *, zoom: bool = False) -> None:
        color = METHOD_COLORS.get(measurement.method, "#666666")
        ours = measurement.method == "SeMARC"
        # Line2D markersize is linear; square-root scaling controls marker area.
        marker_size = (
            (13.8 if ours else 6.0)
            if zoom
            else (16.5 if ours else 8.8) * math.sqrt(MAIN_MARKER_AREA_SCALE)
        )
        ax.plot(
            measurement.latency_s,
            measurement.energy_j,
            linestyle="none",
            marker=METHOD_MARKERS.get(measurement.method, "o"),
            markersize=marker_size,
            markerfacecolor=color,
            markeredgecolor="white",
            markeredgewidth=1.0 if not zoom else 0.7,
            alpha=1.0 if ours else 0.95,
            zorder=7 if ours else 5,
        )

    def set_axis_range(ax, x_values, y_values, *, inset: bool = False) -> None:
        padding = 0.48 if inset else 0.18
        if scale == "log":
            ax.set_xscale("log")
            ax.set_yscale("log")
            x_log = [math.log10(value) for value in x_values]
            y_log = [math.log10(value) for value in y_values]
            x_pad = padding * max(max(x_log) - min(x_log), 0.25)
            y_pad = padding * max(max(y_log) - min(y_log), 0.25)
            ax.set_xlim(
                10 ** (min(x_log) - x_pad), 10 ** (max(x_log) + x_pad)
            )
            ax.set_ylim(
                10 ** (min(y_log) - y_pad), 10 ** (max(y_log) + y_pad)
            )
            ax.xaxis.set_major_locator(
                LogLocator(base=10, subs=(1, 2, 3, 5), numticks=12)
            )
            ax.yaxis.set_major_locator(
                LogLocator(base=10, subs=(1, 2, 3, 5), numticks=12)
            )
            ax.xaxis.set_major_formatter(FuncFormatter(compact_tick))
            ax.yaxis.set_major_formatter(FuncFormatter(compact_tick))
            ax.xaxis.set_minor_formatter(NullFormatter())
            ax.yaxis.set_minor_formatter(NullFormatter())
        else:
            x_span = max(x_values) - min(x_values)
            y_span = max(y_values) - min(y_values)
            x_pad = padding * (x_span if x_span > 0 else max(x_values) * 0.1)
            y_pad = padding * (y_span if y_span > 0 else max(y_values) * 0.1)
            ax.set_xlim(max(0.0, min(x_values) - x_pad), max(x_values) + x_pad)
            ax.set_ylim(max(0.0, min(y_values) - y_pad), max(y_values) + y_pad)
            ax.xaxis.set_major_locator(MaxNLocator(nbins=3 if inset else 4))
            ax.yaxis.set_major_locator(MaxNLocator(nbins=3 if inset else 4))

    for panel_index, (ax, platform) in enumerate(zip(axes.flat, platforms)):
        platform_values = [
            measurement
            for measurement in panel_measurements
            if measurement.platform == platform
        ]
        for measurement in platform_values:
            draw_point(ax, measurement)

        x_values = [measurement.latency_s for measurement in platform_values]
        y_values = [measurement.energy_j for measurement in platform_values]
        set_axis_range(ax, x_values, y_values)

        semarc = next(
            measurement
            for measurement in platform_values
            if measurement.method == "SeMARC"
        )
        ax.annotate(
            f1_label(f1_scores["SeMARC"]),
            xy=(semarc.latency_s, semarc.energy_j),
            xytext=(6, 22),
            textcoords="offset points",
            ha="center",
            va="center",
            fontsize=publication_font(10.0),
            color=METHOD_COLORS["SeMARC"],
            arrowprops={
                "arrowstyle": "-",
                "color": METHOD_COLORS["SeMARC"],
                "linewidth": 0.55,
                "shrinkA": 1.5,
                "shrinkB": 11.0,
            },
            zorder=9,
        )

        competitors = [
            measurement
            for measurement in platform_values
            if measurement.method != "SeMARC"
        ]
        zoom_ax = ax.inset_axes((0.37, 0.06, 0.60, 0.43))
        for measurement in competitors:
            draw_point(zoom_ax, measurement, zoom=True)
        set_axis_range(
            zoom_ax,
            [measurement.latency_s for measurement in competitors],
            [measurement.energy_j for measurement in competitors],
            inset=True,
        )
        if platform == "Android (INT8)" and scale == "linear":
            # Less vertical padding makes this a magnification on both axes,
            # while keeping the original main-panel limits and all data intact.
            energy_min = min(measurement.energy_j for measurement in competitors)
            energy_max = max(measurement.energy_j for measurement in competitors)
            energy_padding = 0.25 * (energy_max - energy_min)
            if energy_padding > 0:
                zoom_ax.set_ylim(energy_min - energy_padding, energy_max + energy_padding)
        positions = annotation_positions.get(platform, {})
        for measurement in competitors:
            color = METHOD_COLORS[measurement.method]
            zoom_ax.annotate(
                f1_label(f1_scores[measurement.method]),
                xy=(measurement.latency_s, measurement.energy_j),
                xytext=positions.get(measurement.method, (0.5, 0.5)),
                textcoords="axes fraction",
                ha="center",
                va="center",
                fontsize=publication_font(8.2),
                color=color,
                arrowprops={
                    "arrowstyle": "-",
                    "color": color,
                    "linewidth": 0.45,
                    "shrinkA": 1.0,
                    "shrinkB": 2.5,
                },
                zorder=9,
            )
        # The source rectangle is exactly the data region shown in the inset.
        # Join corresponding lower source corners to the upper inset corners.
        zoom_x0, zoom_x1 = zoom_ax.get_xlim()
        zoom_y0, zoom_y1 = zoom_ax.get_ylim()
        source_box = Rectangle(
            (zoom_x0, zoom_y0), zoom_x1 - zoom_x0, zoom_y1 - zoom_y0,
            transform=ax.transData, fill=False, edgecolor="black",
            linewidth=1.0, linestyle=":", zorder=4,
        )
        source_box.set_gid(f"zoom-source-{panel_index}")
        ax.add_patch(source_box)
        for side, source_x in ((0, zoom_x0), (1, zoom_x1)):
            connector = ConnectionPatch(
                xyA=(source_x, zoom_y0), coordsA=ax.transData,
                xyB=(side, 1), coordsB=zoom_ax.transAxes,
                color="black", linewidth=0.85, linestyle=":",
                clip_on=False, zorder=3,
            )
            connector.set_gid(f"zoom-connector-{panel_index}-{side}")
            ax.add_artist(connector)
        # Two inset ticks per axis retain a quantitative reading without using
        # extra space between the inset and its parent axes.
        inset_x_ticks = sorted({min(p.latency_s for p in competitors), max(p.latency_s for p in competitors)})
        inset_y_ticks = sorted({min(p.energy_j for p in competitors), max(p.energy_j for p in competitors)})
        tick_text = lambda value: f"{value:.0f}" if value >= 100 else f"{value:.3g}"
        zoom_ax.set_xticks(inset_x_ticks, [tick_text(value) for value in inset_x_ticks])
        zoom_ax.set_yticks(inset_y_ticks, [tick_text(value) for value in inset_y_ticks])
        zoom_ax.tick_params(
            axis="x", direction="out", length=2, width=0.45, pad=1.5,
            labelsize=publication_font(6.2), colors="black",
        )
        zoom_ax.tick_params(
            axis="y", direction="in", length=2, width=0.45, pad=-3,
            labelsize=publication_font(6.2), colors="black",
        )
        for label in zoom_ax.get_yticklabels():
            label.set_horizontalalignment("left")
        zoom_ax.grid(False)
        for spine in zoom_ax.spines.values():
            spine.set_color("black")
            spine.set_linewidth(0.9)

        ax.grid(
            True,
            which="major",
            color="#D8D8D8",
            linewidth=0.55,
            alpha=0.82,
        )
        if scale == "log":
            ax.grid(
                True,
                which="minor",
                color="#ECECEC",
                linewidth=0.35,
                alpha=0.62,
            )
        ax.tick_params(
            which="major", direction="out", length=3.0, width=0.65, pad=2.0
        )
        ax.tick_params(which="minor", direction="out", length=1.8, width=0.45)
        panel_label = f"({chr(ord('a') + panel_index)}) {platform}"
        ax.text(
            0.035,
            0.955,
            panel_label,
            transform=ax.transAxes,
            ha="left",
            va="top",
            fontsize=publication_font(11.5),
            color="black",
        )
        for spine in ax.spines.values():
            spine.set_color("black")
            spine.set_linewidth(1.0)

    method_handles = [
        Line2D(
            [0],
            [0],
            linestyle="none",
            marker=METHOD_MARKERS.get(method, "o"),
            markerfacecolor=METHOD_COLORS.get(method, "#666666"),
            markeredgecolor="white",
            markeredgewidth=0.8,
            markersize=11.0 if method == "SeMARC" else 7.3,
            label=METHOD_LABELS.get(method, method),
        )
        for method in methods
    ]
    legend = fig.legend(
        handles=method_handles,
        loc="upper center",
        bbox_to_anchor=(0.5, 0.995),
        ncol=len(method_handles),
        frameon=True,
        fancybox=False,
        framealpha=1.0,
        facecolor="white",
        edgecolor="#B8B8B8",
        columnspacing=0.45,
        handlelength=0.65,
        handletextpad=0.20,
        borderpad=0.20,
        labelspacing=0.35,
        fontsize=publication_font(8.6),
    )
    legend.get_frame().set_linewidth(0.55)

    scale_suffix = ", log scale" if scale == "log" else ""
    fig.supxlabel(f"Latency per inference (s{scale_suffix})", y=0.012)
    fig.supylabel(f"Energy per inference (J{scale_suffix})", x=0.012)
    fig.subplots_adjust(
        left=0.072,
        right=0.99,
        bottom=0.148,
        top=0.88,
        wspace=0.20,
    )
    # Keep the requested font sizes and complete model labels in one row.
    fig.canvas.draw()
    legend_width = legend.get_window_extent(fig.canvas.get_renderer()).width / fig.dpi
    if legend_width + 0.30 > fig.get_figwidth():
        fig.set_size_inches(legend_width + 0.30, fig.get_figheight())

    paths = save_figure(fig, output_dir, stem, dpi)
    print(f"Rendered {dataset_name}: {', '.join(platforms)} ({scale} axes).")
    return paths


def output_stem(dataset: str, scale: str) -> str:
    if dataset.casefold() == "all":
        return "energy_latency_full"
    slug = re.sub(r"[^a-z0-9]+", "_", dataset.casefold()).strip("_")
    return f"energy_latency_{slug}_3panel_{scale}"


def main() -> None:
    args = parse_args()
    book = load_workbook(args.input)
    measurements, empty_blocks, incomplete_pairs = read_measurements(
        book, args.sheet
    )
    stem = args.stem or output_stem(args.dataset, args.scale)
    if args.dataset.casefold() == "all":
        paths = render(measurements, args.output_dir, stem, args.dpi)
    else:
        f1_scores = read_f1_scores(
            book, args.results_sheet, args.dataset
        )
        paths = render_dataset_platforms(
            measurements,
            args.dataset,
            args.scale,
            f1_scores,
            args.output_dir,
            stem,
            args.dpi,
        )

    print(f"Read {len(measurements)} populated latency/energy pairs.")
    for label, values in (
        ("Platforms", {measurement.platform for measurement in measurements}),
        ("Datasets", {measurement.dataset for measurement in measurements}),
        ("Models", {measurement.method for measurement in measurements}),
    ):
        print(f"{label}: {', '.join(sorted(values))}")
    if empty_blocks:
        print(
            "Skipped empty platform blocks: "
            + ", ".join(empty_blocks)
        )
    for warning in incomplete_pairs:
        print(f"Warning: {warning}")
    print("Created:")
    for path in paths:
        print(f"  {path.resolve()}")


if __name__ == "__main__":
    main()
