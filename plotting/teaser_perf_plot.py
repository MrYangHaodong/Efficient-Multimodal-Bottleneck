#!/usr/bin/env python3
"""Plot absolute IEMOCAP Android INT8 latency and energy with Macro-F1.

The figure compares full-availability measurements for SeMARC, MBT, AdaMML,
and DyMo on Android INT8. Raw latency is converted from ms to s and energy is
converted from mJ to J.

Dependencies: matplotlib, numpy, openpyxl
"""

from __future__ import annotations

import argparse
import base64
import urllib.request
from io import BytesIO
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
from matplotlib.ticker import MaxNLocator


SHEET_URL = (
    "https://docs.google.com/spreadsheets/d/"
    "1ZsKgZztYMPOj76LL7eYWT07JJ4k0cWdaC1zmOlIy3q0/"
    "export?format=xlsx"
)

METHODS = ["SeMARC", "MBT", "AdaMML", "DyMo"]
LEGEND_LABELS = {
    "SeMARC": "SeMARC (Ours)",
    "MBT": "MBT (Monolithic Fusion)",
    "AdaMML": "AdaMML (Selective Fusion)",
    "DyMo": "DyMo (Selective Fusion)",
}
COLORS = {
    "SeMARC": "#AA4499",  # Paul Tol purple
    "MBT": "#EE7733",  # Paul Tol orange
    "AdaMML": "#4477AA",  # Paul Tol blue
    "DyMo": "#882255",  # Paul Tol wine/maroon
}
MARKERS = {
    "SeMARC": "*",  # proactive sequential acquisition
    "MBT": "o",  # monolithic fusion
    "AdaMML": "s",  # selective fusion
    "DyMo": "s",  # selective fusion
}

SHEET_METHODS = {
    "Ours": "SeMARC",
    "Vanilla MBT": "MBT",
    "Vanilla MBT (est.)": "MBT",
    "AdaMML": "AdaMML",
    "DyMo": "DyMo",
}
F1_ROWS = {
    "SeMARC": "Ours",
    "MBT": "Vanilla MBT (R)",
    "AdaMML": "AdaMML (E)",
    "DyMo": "DyMo (E)",
}

# platform: anchor label, variant label, latency scale, energy scale
PLATFORMS = {
    "Android (INT8)": ("Onnx_Android", "Quantized (INT8)", 1e-3, 1e-3),
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--input",
        default=SHEET_URL,
        help="Local workbook path or Google Sheets XLSX export URL.",
    )
    parser.add_argument("--dataset", default="IEMOCAP")
    parser.add_argument("--energy-sheet", default="Energy_Study")
    parser.add_argument("--results-sheet", default="Main_ResultsFull")
    parser.add_argument(
        "--output-dir", type=Path, default=Path("output/figures")
    )
    parser.add_argument(
        "--stem", default="iemocap_android_energy_latency_accuracy_full"
    )
    parser.add_argument("--inline-html", type=Path, default=None)
    return parser.parse_args()


def load_workbook(source: str):
    if source.startswith(("http://", "https://")):
        request = urllib.request.Request(
            source, headers={"User-Agent": "Mozilla/5.0"}
        )
        with urllib.request.urlopen(request, timeout=60) as response:
            return openpyxl.load_workbook(
                BytesIO(response.read()), data_only=True
            )
    return openpyxl.load_workbook(
        Path(source).expanduser(), data_only=True
    )


def first_number(value) -> float:
    if isinstance(value, (int, float)):
        return float(value)
    return float(str(value).split("|")[0].strip())


def find_block(ws, anchor: str, variant: str | None) -> tuple[int, int]:
    for row in range(1, ws.max_row + 1):
        first = str(ws.cell(row, 1).value or "").strip()
        second = str(ws.cell(row, 2).value or "").strip()
        if first == anchor and (variant is None or second == variant):
            method_row = row + 1
            data_row = next(
                candidate
                for candidate in range(method_row + 2, ws.max_row + 1)
                if str(ws.cell(candidate, 1).value or "").strip()
                == "IEMOCAP"
            )
            return method_row, data_row
    raise KeyError(f"Could not find platform block: {anchor}, {variant}")


def read_measurements(book, sheet_name: str, dataset: str):
    ws = book[sheet_name]
    measurements = {platform: {} for platform in PLATFORMS}

    for platform, (anchor, variant, latency_scale, energy_scale) in PLATFORMS.items():
        method_row, first_data_row = find_block(ws, anchor, variant)
        data_row = next(
            row
            for row in range(first_data_row, ws.max_row + 1)
            if str(ws.cell(row, 1).value or "").strip() == dataset
        )
        for column in range(2, ws.max_column + 1):
            raw_method = str(ws.cell(method_row, column).value or "").strip()
            method = SHEET_METHODS.get(raw_method)
            if method not in METHODS:
                continue
            latency = first_number(ws.cell(data_row, column - 1).value)
            energy = first_number(ws.cell(data_row, column).value)
            measurements[platform][method] = {
                "latency_s": latency * latency_scale,
                "energy_j": energy * energy_scale,
            }

        missing = set(METHODS) - set(measurements[platform])
        if missing:
            raise ValueError(
                f"Missing {platform} measurements for: {sorted(missing)}"
            )
    return measurements


def read_accuracy(book, sheet_name: str):
    ws = book[sheet_name]
    row_by_name = {
        str(ws.cell(row, 1).value or "").strip(): row
        for row in range(1, ws.max_row + 1)
    }
    return {
        method: first_number(ws.cell(row_by_name[row_name], 2).value)
        for method, row_name in F1_ROWS.items()
    }


def annotation_offsets():
    return {
        "SeMARC": (7, 7),
        "MBT": (-8, 17),
        "AdaMML": (-8, -21),
        "DyMo": (7, 8),
    }


def render(measurements, accuracy, output_dir: Path, stem: str):
    plt.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "font.size": 16.6,
            "axes.labelsize": 17.4,
            "xtick.labelsize": 16.1,
            "ytick.labelsize": 16.1,
            "legend.fontsize": 13.9,
            "axes.linewidth": 0.7,
            "pdf.fonttype": 42,
            "ps.fonttype": 42,
            "svg.fonttype": "none",
        }
    )

    fig, ax = plt.subplots(figsize=(5.0, 4.0))
    platform_values = measurements["Android (INT8)"]
    latencies = np.array(
        [platform_values[method]["latency_s"] for method in METHODS]
    )
    energies = np.array(
        [platform_values[method]["energy_j"] for method in METHODS]
    )

    offsets = annotation_offsets()
    for method in METHODS:
        x = platform_values[method]["latency_s"]
        y = platform_values[method]["energy_j"]
        ours = method == "SeMARC"
        ax.scatter(
            x,
            y,
            s=650 if ours else 78,
            marker=MARKERS[method],
            facecolor=COLORS[method],
            edgecolor="white" if ours else COLORS[method],
            linewidth=1.6 if ours else 0.8,
            zorder=6 if ours else 4,
            label=LEGEND_LABELS[method],
        )

        dx, dy = offsets[method]
        ax.annotate(
            f"F1={accuracy[method]:.2f}",
            xy=(x, y),
            xytext=(dx, dy),
            textcoords="offset points",
            ha="right" if dx < 0 else "left",
            va="center",
            fontsize=15.7,
            color=COLORS[method],
            arrowprops={
                "arrowstyle": "-",
                "color": COLORS[method],
                "linewidth": 0.5,
                "shrinkA": 1.5,
                "shrinkB": 3.5,
            },
            zorder=8,
        )

    x_span = latencies.max() - latencies.min()
    y_span = energies.max() - energies.min()
    ax.set_xlim(
        max(0.0, latencies.min() - 0.18 * x_span),
        latencies.max() + 0.42 * x_span,
    )
    ax.set_ylim(
        max(0.0, energies.min() - 0.18 * y_span),
        energies.max() + 0.22 * y_span,
    )
    ax.xaxis.set_major_locator(MaxNLocator(nbins=4))
    ax.yaxis.set_major_locator(MaxNLocator(nbins=4))
    ax.grid(True, color="#DEDEDE", linewidth=0.45, alpha=0.78)
    ax.tick_params(direction="out", length=2.2, width=0.55, pad=1.5)
    ax.set_xlabel("Latency per inference (s)", fontsize=17.8)
    ax.set_ylabel("Energy per inference (J)", fontsize=17.8)
    ax.yaxis.set_label_coords(-0.12, 0.46)
    for spine in ax.spines.values():
        spine.set_visible(True)
        spine.set_color("#444444")
        spine.set_linewidth(0.7)

    handles, labels = ax.get_legend_handles_labels()
    legend = fig.legend(
        handles,
        labels,
        loc="upper center",
        bbox_to_anchor=(0.5, 1.035),
        ncol=2,
        frameon=True,
        fancybox=False,
        framealpha=1.0,
        facecolor="#FAFAFA",
        edgecolor="#D0D0D0",
        columnspacing=0.45,
        handlelength=0.7,
        handletextpad=0.15,
        borderpad=0.15,
        labelspacing=0.2,
        markerscale=0.85,
    )
    legend.get_frame().set_linewidth(0.55)
    fig.subplots_adjust(
        left=0.13,
        right=0.995,
        bottom=0.15,
        top=0.865,
    )

    output_dir.mkdir(parents=True, exist_ok=True)
    paths = []
    for suffix, options in (
        ("pdf", {}),
        ("svg", {}),
        ("png", {"dpi": 400}),
    ):
        path = output_dir / f"{stem}.{suffix}"
        fig.savefig(path, bbox_inches="tight", pad_inches=0.005, **options)
        paths.append(path)
    plt.close(fig)
    return paths


def write_inline_fragment(svg_path: Path, destination: Path) -> None:
    template_path = Path(__file__).parent / "templates" / "inline-svg-figure.html"
    template = template_path.read_text(encoding="utf-8")
    encoded = base64.b64encode(svg_path.read_bytes()).decode("ascii")
    fragment = (
        template.replace("__ROOT_ID__", "iemocap-energy-latency-accuracy")
        .replace("__SVG_DATA__", encoded)
        .replace(
            "__ALT_TEXT__",
            "Absolute IEMOCAP energy and latency on Android INT8; point labels report Macro-F1.",
        )
    )
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(fragment, encoding="utf-8")


def main() -> None:
    args = parse_args()
    book = load_workbook(args.input)
    measurements = read_measurements(
        book, args.energy_sheet, args.dataset
    )
    accuracy = read_accuracy(book, args.results_sheet)
    paths = render(measurements, accuracy, args.output_dir, args.stem)
    if args.inline_html is not None:
        svg_path = next(path for path in paths if path.suffix == ".svg")
        write_inline_fragment(svg_path, args.inline_html)
        paths.append(args.inline_html)

    for platform, values in measurements.items():
        print(platform)
        for method in METHODS:
            point = values[method]
            print(
                f"  {method}: {point['latency_s']:.5g} s, "
                f"{point['energy_j']:.5g} J, F1={accuracy[method]:.3f}"
            )
    print("Created:")
    for path in paths:
        print(f"  {path.resolve()}")


if __name__ == "__main__":
    main()
